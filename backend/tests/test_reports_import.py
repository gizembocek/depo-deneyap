import io

from openpyxl import Workbook, load_workbook
from pypdf import PdfReader


def _pdf_text(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    return "\n".join(page.extract_text() for page in reader.pages)


def test_pdf_turkce_karakterler(client, admin, product):
    """Eski sürümde ş, ğ, ı harfleri ■ olarak çıkıyordu."""
    product("Işıklı Büyüteç Şarjlı", 1)
    product("Çağdaş Öğretim Seti", 50)
    r = client.get("/api/reports/export/pdf", headers=admin)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    text = _pdf_text(r.content)
    assert "Işıklı Büyüteç Şarjlı" in text
    assert "Çağdaş Öğretim Seti" in text
    assert "Artvin Deneyap Atölyesi" in text
    assert "■" not in text


def test_pdf_cok_sayfali_ve_bos(client, admin, product):
    for i in range(120):
        product(f"Malzeme {i} - uzun bir ürün adı ile tablonun satır kaydırmasını deniyoruz", i)
    r = client.get("/api/reports/export/pdf", headers=admin)
    assert len(PdfReader(io.BytesIO(r.content)).pages) > 2
    empty = client.get("/api/reports/export/pdf", params={"category": "Sarf Malzemesi"}, headers=admin)
    assert "eşleşen ürün yok" in _pdf_text(empty.content)


def test_excel_raporu_ve_filtre(client, admin, product, warehouse):
    w = warehouse("Elektronik Depo")
    product("Direnç Seti", 2, warehouse_id=w)
    product("Kontrplak", 40, category="Sarf Malzemesi")
    r = client.get("/api/reports/export/excel", params={"warehouse_id": w}, headers=admin)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Ürün Raporu"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1 and rows[0][1] == "Direnç Seti" and rows[0][5] == "Elektronik Depo"


def test_rapor_eski_token_parametresi_hala_calisir(client, admin):
    token = admin["Authorization"].split()[1]
    assert client.get(f"/api/reports/export/excel?token={token}").status_code == 200
    assert client.get("/api/reports/export/excel").status_code == 401
    assert client.get("/api/reports/export/pdf?token=bozuk").status_code == 401


def test_rapor_kullanici_yetkisine_gore(client, admin, make_user, warehouse, product):
    w1, w2 = warehouse("A"), warehouse("B")
    product("Görünür", warehouse_id=w1)
    product("Görünmez", warehouse_id=w2)
    user_h, _ = make_user(warehouse_ids=[w1])
    text = _pdf_text(client.get("/api/reports/export/pdf", headers=user_h).content)
    assert "Görünür" in text and "Görünmez" not in text


def test_islem_kayitlari(client, admin, product):
    p = product("Lehim", 5)
    client.post(f"/api/products/{p['id']}/stock", headers=admin, json={"action": "add", "quantity": 2, "note": "Bağış"})
    logs = client.get("/api/reports/logs?limit=10", headers=admin).json()
    assert logs[0]["action"] == "add" and logs[0]["note"] == "Bağış"
    assert logs[0]["created_at"].endswith("+03:00")
    assert client.get("/api/reports/logs?limit=999999", headers=admin).status_code == 422


# ---------------------------------------------------------------------------
# Excel içe aktarma
# ---------------------------------------------------------------------------
def _xlsx(sheets: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _upload(client, admin, content, name="liste.xlsx", warehouse_id=None):
    data = {"warehouse_id": str(warehouse_id)} if warehouse_id else {}
    return client.post("/api/products/import/excel", headers=admin, data=data,
                       files={"file": (name, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def test_deneyap_envanter_formati(client, admin, warehouse):
    """Atölyelerdeki tipik envanter dosyası: başlık 3. satırda, miktar 'Birim' içinde."""
    w = warehouse()
    content = _xlsx({
        "Robotik ve Kodlama": [
            ["ARTVİN DENEYAP TEKNOLOJİ ATÖLYESİ"], [],
            ["Sıra", "Malzeme Adı", "Birim", "Çalışır Durumda mı?"],
            [1, "Arduino Uno", "12 Adet", "Evet"],
            [2, "Servo Motor SG90", "20 adet", "Kısmen"],
            [3, "Mesafe Sensörü", "Adet", "Bozuk"],
            [None, "TOPLAM", None, None],
        ],
        "Demirbaş Listesi": [
            ["Malzeme Adı", "Miktar", "Açıklama"],
            ["3D Yazıcı", 2, "Ender 3"],
            ["Arduino Uno", 1, "tekrar eden satır"],
        ],
        "DURUM": [["Toplam puan", 5]],
    })
    r = _upload(client, admin, content, warehouse_id=w)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["added"] == 4 and body["skipped"] == 1
    assert "DURUM" in body["skipped_sheets"]

    products = {p["name"]: p for p in client.get("/api/products/?per_page=200", headers=admin).json()}
    assert products["Arduino Uno"]["current_stock"] == 12
    assert products["Servo Motor SG90"]["status"] == "Garantide"
    assert products["Mesafe Sensörü"]["status"] == "Bozuk / Kırık"
    assert products["Mesafe Sensörü"]["current_stock"] == 0
    assert products["3D Yazıcı"]["description"] == "Ender 3"
    # 'Robotik ve Kodlama' sayfası mevcut 'Robotik Kodlama' dersiyle eşleşir (yeni ders açılmaz)
    assert [c["name"] for c in products["Arduino Uno"]["courses"]] == ["Demirbaş Listesi", "Robotik Kodlama"]
    course_names = [c["name"] for c in client.get("/api/courses/", headers=admin).json()]
    assert "Robotik ve Kodlama" not in course_names


def test_sade_liste_ve_genel_sayfa_adi(client, admin):
    content = _xlsx({"Sayfa1": [["Ürün Adı"], ["Makas"], ["Cetvel"], ["Makas"]]})
    body = _upload(client, admin, content).json()
    assert body["added"] == 2
    # 'Sayfa1' ders olarak eklenmez
    assert all(c["name"] != "Sayfa1" for c in client.get("/api/courses/", headers=admin).json())


def test_kategori_durum_kritik_ders_sutunlari(client, admin):
    content = _xlsx({"Liste": [
        ["Malzeme", "Adet", "Kategori", "Durum", "Kritik Stok", "İdeal Stok", "Ders"],
        ["Filament", "1.250", "Sarf", "çalışıyor", 3, 20, "Tasarım ve Üretim, Robotik Kodlama"],
    ]})
    assert _upload(client, admin, content).status_code == 200
    p = client.get("/api/products/", headers=admin).json()[0]
    assert (p["current_stock"], p["category"], p["status"], p["critical_stock"], p["ideal_stock"]) == \
           (1250, "Sarf Malzemesi", "Çalışan", 3, 20)
    assert sorted(c["name"] for c in p["courses"]) == ["Robotik Kodlama", "Tasarım ve Üretim"]


def test_sablon_indir_ve_geri_yukle(client, admin):
    r = client.get("/api/products/import/template", headers=admin)
    assert r.status_code == 200
    body = _upload(client, admin, r.content, name="sablon.xlsx").json()
    assert body["added"] == 3


def test_gecersiz_dosyalar(client, admin):
    assert _upload(client, admin, b"x", name="eski.xls").status_code == 400
    assert "xlsx" in _upload(client, admin, b"x", name="eski.xls").json()["detail"]
    assert _upload(client, admin, b"bozuk icerik", name="bozuk.xlsx").status_code == 400
    assert _upload(client, admin, b"a,b", name="liste.csv").status_code == 400
    r = _upload(client, admin, _xlsx({"A": [["Fiyat", "Tarih"], [1, 2]]}))
    assert r.status_code == 400 and "Malzeme Adı" in r.json()["detail"]
