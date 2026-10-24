"""
Excel'den ürün aktarma.

Atölyelerin elindeki envanter listeleri birbirinden farklıdır. Bu modül:
  - Tüm sayfaları tarar, ilk 20 satırda başlık satırını otomatik bulur.
  - Zorunlu sütun yalnızca ürün adıdır ("Malzeme Adı", "Ürün Adı", "Malzeme", "Ad" ...).
  - Miktar "Miktar/Adet/Stok" sütunundan; yoksa "Birim" sütunundaki sayıdan ("6 Adet" -> 6) okunur.
  - Kategori, durum, açıklama, kritik stok, ideal stok ve ders sütunları isteğe bağlıdır.
  - Sayfa adı bir derse karşılık geliyorsa ürünler o derse bağlanır.

Bu modül veritabanına dokunmaz; sadece Excel'i okuyup satır listesi döndürür.
"""
import io
import re
from dataclasses import dataclass, field
from typing import List, Optional

from database import tr_fold
from models import CATEGORY_VALUES

GENERIC_SHEET_NAMES = {
    "sheet", "sheet1", "sheet2", "sheet3", "sayfa", "sayfa1", "sayfa2", "sayfa3",
    "liste", "urunler", "malzemeler", "malzeme listesi", "envanter", "stok", "depo",
}

SKIP_NAMES = {"", "none", "nan", "toplam", "genel toplam", "ara toplam", "-", "—"}


def norm(text) -> str:
    """Başlıkları karşılaştırmak için sadeleştirir: 'Malzeme Adı:' -> 'malzeme adi'."""
    if text is None:
        return ""
    text = tr_fold(str(text))
    text = re.sub(r"[^a-z0-9 ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


NAME_EXACT = {
    "malzeme adi", "malzemenin adi", "malzeme ismi", "malzeme", "malzeme cinsi",
    "urun adi", "urunun adi", "urun ismi", "urun", "ad", "adi", "isim",
    "material", "material name", "name", "product", "product name", "item",
}
QTY_EXACT = {
    "miktar", "miktari", "adet", "adedi", "stok", "stok miktari", "stok adedi", "mevcut",
    "mevcut stok", "mevcut miktar", "mevcut adet", "sayi", "sayisi", "toplam", "toplam adet",
    "quantity", "qty", "stock", "count",
}
CATEGORY_EXACT = {"kategori", "kategorisi", "tur", "turu", "tip", "tipi", "cinsi", "category", "malzeme turu"}
DESC_EXACT = {"aciklama", "aciklamasi", "ozellik", "ozellikler", "ozellikleri", "not", "notlar", "description", "detay"}
COURSE_EXACT = {"ders", "dersler", "ders adi", "kullanildigi ders", "kullanilan ders", "course"}


@dataclass
class ColumnMap:
    name: int = -1
    qty: int = -1
    birim: int = -1
    category: int = -1
    status: int = -1
    description: int = -1
    critical: int = -1
    ideal: int = -1
    course: int = -1


def detect_columns(headers: List[str]) -> ColumnMap:
    cols = ColumnMap()
    # Uzun metinler (açıklama cümleleri) başlık sayılmaz
    normalized = [norm(h) if len(str(h or "")) <= 40 else "" for h in headers]

    # 1) Tam eşleşmeler (en güvenilir)
    for i, h in enumerate(normalized):
        if not h:
            continue
        if cols.name < 0 and h in NAME_EXACT:
            cols.name = i
        elif cols.qty < 0 and h in QTY_EXACT:
            cols.qty = i
        elif cols.category < 0 and h in CATEGORY_EXACT:
            cols.category = i
        elif cols.description < 0 and h in DESC_EXACT:
            cols.description = i
        elif cols.course < 0 and h in COURSE_EXACT:
            cols.course = i

    # 2) İçeren eşleşmeler
    for i, h in enumerate(normalized):
        if not h or i in (cols.name, cols.qty, cols.category, cols.description, cols.course):
            continue
        if cols.name < 0 and ("malzeme adi" in h or "urun adi" in h or "malzeme ismi" in h or "urun ismi" in h):
            cols.name = i
        elif cols.birim < 0 and "birim" in h:
            cols.birim = i
        elif cols.critical < 0 and "kritik" in h:
            cols.critical = i
        elif cols.ideal < 0 and ("ideal" in h or "hedef" in h or "olmasi gereken" in h):
            cols.ideal = i
        elif cols.status < 0 and ("durum" in h or "calis" in h or "status" in h or "saglam" in h):
            cols.status = i
        elif cols.qty < 0 and ("miktar" in h or "adet" in h or ("stok" in h and "kod" not in h)):
            cols.qty = i
        elif cols.description < 0 and ("aciklama" in h or "ozellik" in h):
            cols.description = i
    return cols


def parse_int(value) -> Optional[int]:
    """'6 Adet' -> 6, 12.0 -> 12, '1.250' -> 1250, 'Adet' -> None."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return max(0, int(value))
    text = str(value).strip()
    match = re.search(r"\d[\d.,]*", text)
    if not match:
        return None
    number = match.group(0)
    if re.fullmatch(r"\d{1,3}([.,]\d{3})+", number):  # binlik ayraç: 1.250
        number = re.sub(r"[.,]", "", number)
    else:
        number = number.replace(",", ".")
    try:
        return max(0, int(float(number)))
    except ValueError:
        return None


def parse_status(value) -> str:
    text = norm(value)
    if not text:
        return "Çalışan"
    words = set(text.split())
    if "demirbas" in text and ("dus" in text or "cik" in text):
        return "Demirbaştan düşülecek"
    if "kirik kutu" in text:
        return "Kırık ürün kutusunda"
    if "bozuk" in text or "kirik" in text or "calismiyor" in text or "arizali" in text or "hasarli" in text:
        return "Bozuk / Kırık"
    if "garanti" in text or "tamir" in text or "servis" in text or "bakim" in text or "kismen" in text or "tadilat" in text:
        return "Garantide"
    if "calis" in text or "saglam" in text or words & {"tam", "evet", "e", "yes", "ok", "var"}:
        return "Çalışan"
    if words & {"hayir", "yok", "h", "no", "0"}:
        return "Bozuk / Kırık"
    return "Çalışan"


def parse_category(value, sheet_name: str) -> str:
    text = norm(value)
    if text:
        for cat in CATEGORY_VALUES:
            if norm(cat) == text:
                return cat
        if "sarf" in text:
            return "Sarf Malzemesi"
        if "elektronik" in text or "bilesen" in text:
            return "Elektronik Bileşen"
        if "demirbas" in text or "dayanikli" in text or "genel" in text:
            return "Genel / Dayanıklı Malzeme"
    return category_from_sheet(sheet_name)


def category_from_sheet(sheet_name: str) -> str:
    name = norm(sheet_name)
    if "demirbas" in name:
        return "Genel / Dayanıklı Malzeme"
    if "elektronik" in name or "programlama" in name or "robotik" in name or "kodlama" in name:
        return "Elektronik Bileşen"
    if "sarf" in name or "tasarim" in name or "uretim" in name:
        return "Sarf Malzemesi"
    return "Genel / Dayanıklı Malzeme"


def course_key(name: str) -> str:
    """'Robotik ve Kodlama' ile 'Robotik Kodlama' aynı sayılsın diye 've' atılır."""
    return " ".join(w for w in norm(name).split() if w != "ve")


def is_generic_sheet(sheet_name: str) -> bool:
    return norm(sheet_name) in GENERIC_SHEET_NAMES


@dataclass
class ParsedRow:
    sheet: str
    row_number: int
    name: str
    current_stock: int = 0
    category: str = "Genel / Dayanıklı Malzeme"
    status: str = "Çalışan"
    description: Optional[str] = None
    critical_stock: Optional[int] = None
    ideal_stock: Optional[int] = None
    course_names: List[str] = field(default_factory=list)


@dataclass
class ParseResult:
    rows: List[ParsedRow] = field(default_factory=list)
    sheets: List[str] = field(default_factory=list)
    skipped_sheets: List[str] = field(default_factory=list)


def _cell(row, index):
    if index < 0 or index >= len(row):
        return None
    return row[index]


def parse_workbook(contents: bytes) -> ParseResult:
    import openpyxl

    wb = openpyxl.load_workbook(filename=io.BytesIO(contents), data_only=True, read_only=True)
    result = ParseResult()
    try:
        for sheet in wb.worksheets:
            header_row_idx = -1
            cols = None
            for idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=True), 1):
                if not row:
                    continue
                candidate = detect_columns([h if h is not None else "" for h in row])
                # Başlık satırı: ad sütunu + en az bir yardımcı sütun (veya tek sütunlu liste)
                helpers = [candidate.qty, candidate.birim, candidate.category, candidate.status,
                           candidate.description, candidate.critical, candidate.ideal, candidate.course]
                non_empty = sum(1 for h in row if h not in (None, ""))
                single_column_list = non_empty == 1 and norm(row[candidate.name] if candidate.name >= 0 else "") in NAME_EXACT
                if candidate.name >= 0 and (any(h >= 0 for h in helpers) or single_column_list):
                    header_row_idx, cols = idx, candidate
                    break

            if cols is None:
                result.skipped_sheets.append(sheet.title)
                continue

            sheet_courses = [] if is_generic_sheet(sheet.title) else [sheet.title]
            count = 0
            for row_number, row in enumerate(sheet.iter_rows(min_row=header_row_idx + 1, values_only=True), header_row_idx + 1):
                raw_name = _cell(row, cols.name)
                if raw_name is None:
                    continue
                name = re.sub(r"\s+", " ", str(raw_name)).strip()
                if norm(name) in SKIP_NAMES or re.fullmatch(r"[\d.,\s]+", name):
                    continue
                if len(name) > 255:
                    name = name[:255]

                qty = parse_int(_cell(row, cols.qty)) if cols.qty >= 0 else None
                if qty is None and cols.birim >= 0:
                    qty = parse_int(_cell(row, cols.birim))

                description = _cell(row, cols.description)
                course_names = list(sheet_courses)
                course_cell = _cell(row, cols.course)
                if course_cell:
                    course_names += [c.strip() for c in re.split(r"[,;/]", str(course_cell)) if c.strip()]

                result.rows.append(ParsedRow(
                    sheet=sheet.title,
                    row_number=row_number,
                    name=name,
                    current_stock=qty or 0,
                    category=parse_category(_cell(row, cols.category), sheet.title),
                    status=parse_status(_cell(row, cols.status)) if cols.status >= 0 else "Çalışan",
                    description=str(description).strip()[:4000] if description not in (None, "") else None,
                    critical_stock=parse_int(_cell(row, cols.critical)) if cols.critical >= 0 else None,
                    ideal_stock=parse_int(_cell(row, cols.ideal)) if cols.ideal >= 0 else None,
                    course_names=course_names,
                ))
                count += 1
            if count:
                result.sheets.append(sheet.title)
            else:
                result.skipped_sheets.append(sheet.title)
    finally:
        wb.close()
    return result


def build_template() -> bytes:
    """Atölyelerin doldurup yükleyebileceği örnek Excel şablonu."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from models import STATUS_VALUES

    wb = Workbook()
    ws = wb.active
    ws.title = "Malzemeler"
    headers = ["Malzeme Adı", "Miktar", "Kategori", "Durum", "Kritik Stok", "İdeal Stok", "Ders", "Açıklama"]
    ws.append(headers)
    examples = [
        ["Arduino Uno R3", 25, "Elektronik Bileşen", "Çalışan", 5, 30, "Robotik Kodlama", "USB kablosu ile"],
        ["PLA Filament 1.75mm", 12, "Sarf Malzemesi", "Çalışan", 3, 20, "Tasarım ve Üretim", "Beyaz renk"],
        ["3 Boyutlu Yazıcı", 2, "Genel / Dayanıklı Malzeme", "Garantide", 0, 2, "Tasarım ve Üretim", ""],
    ]
    for r in examples:
        ws.append(r)
    fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    widths = [32, 10, 28, 24, 12, 12, 28, 30]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"

    dv_cat = DataValidation(type="list", formula1='"' + ",".join(CATEGORY_VALUES) + '"', allow_blank=True)
    dv_status = DataValidation(type="list", formula1='"' + ",".join(STATUS_VALUES) + '"', allow_blank=True)
    ws.add_data_validation(dv_cat)
    ws.add_data_validation(dv_status)
    dv_cat.add("C2:C2000")
    dv_status.add("D2:D2000")

    info = wb.create_sheet("Nasıl Doldurulur")
    lines = [
        "Deneyap Depo Sistemi - Excel İçe Aktarma Şablonu",
        "",
        "• Zorunlu tek sütun 'Malzeme Adı'dır. Diğer sütunlar boş bırakılabilir.",
        "• Miktar boşsa 0 kabul edilir. '6 Adet' gibi yazımlar da okunur.",
        "• Kategori: Sarf Malzemesi / Genel / Dayanıklı Malzeme / Elektronik Bileşen",
        "• Durum: " + " / ".join(STATUS_VALUES),
        "• Ders sütununa birden fazla ders virgülle yazılabilir.",
        "• Birden fazla sayfa kullanabilirsiniz. Sayfa adı bir ders adıysa ürünler o derse bağlanır.",
        "• Seçilen depoda aynı isimde ürün varsa tekrar eklenmez (atlanır).",
    ]
    for line in lines:
        info.append([line])
    info["A1"].font = Font(bold=True, size=14)
    info.column_dimensions["A"].width = 100

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
