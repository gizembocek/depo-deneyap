import io
import os
from datetime import datetime, timedelta, timezone
from typing import Optional

from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload, selectinload

from auth import apply_product_access, user_from_token
from database import get_db, tr_ordered
from models import Product, StockLog, User
from routers.products import apply_product_filters
from settings_store import get_settings

TZ_TR = timezone(timedelta(hours=3))
ASSETS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "assets")

router = APIRouter(prefix="/api/reports", tags=["Reports"])

ACTION_LABELS = {
    "add": "Stok Ekleme", "remove": "Stok Çıkarma", "set": "Stok Ayarlama",
    "create": "Ürün Oluşturma", "status_change": "Durum Değişikliği",
}


def to_tr_time(dt):
    """UTC datetime'ı Türkiye saatine (UTC+3) çevirir."""
    if dt is None:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(TZ_TR).isoformat()


def _tr_now() -> datetime:
    return datetime.now(timezone.utc).astimezone(TZ_TR)


def get_report_user(request: Request, token: Optional[str] = Query(None), db: Session = Depends(get_db)) -> User:
    """Önce Authorization başlığına, yoksa (eski sürüm uyumluluğu için) ?token= parametresine bakar."""
    header = request.headers.get("authorization", "")
    if header.lower().startswith("bearer "):
        return user_from_token(header[7:].strip(), db)
    return user_from_token(token, db)


def _report_products(db: Session, user: User, category, warehouse_id, status, low_stock, course_id):
    query = db.query(Product).options(selectinload(Product.courses), joinedload(Product.warehouse),
                                      joinedload(Product.cabinet), joinedload(Product.shelf),
                                      joinedload(Product.compartment))
    query = apply_product_filters(apply_product_access(query, user), category=category, status=status,
                                  warehouse_id=warehouse_id, low_stock=low_stock, course_id=course_id)
    return query.order_by(tr_ordered(Product.name), Product.id).all()


def _location(p: Product) -> str:
    parts = [p.cabinet.name if p.cabinet else None, p.shelf.name if p.shelf else None,
             p.compartment.name if p.compartment else None]
    return " / ".join(x for x in parts if x)


def _title(db: Session) -> str:
    atolye = get_settings(db).get("atolye_adi")
    return f"{atolye} - Stok Raporu" if atolye else "Deneyap Atölyesi - Stok Raporu"


@router.get("/logs")
def get_all_logs(
    product_id: Optional[int] = None,
    user_id: Optional[int] = None,
    action: Optional[str] = None,
    limit: int = Query(100, ge=1, le=1000),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_report_user),
):
    query = db.query(StockLog).options(joinedload(StockLog.user), joinedload(StockLog.product))
    if product_id:
        query = query.filter(StockLog.product_id == product_id)
    if user_id:
        query = query.filter(StockLog.user_id == user_id)
    if action:
        query = query.filter(StockLog.action == action)
    if current_user.role != "admin":
        query = apply_product_access(query.join(StockLog.product), current_user)

    logs = query.order_by(StockLog.created_at.desc(), StockLog.id.desc()).limit(limit).all()
    return [{
        "id": log.id,
        "product_id": log.product_id,
        "product_name": log.product.name if log.product else None,
        "user_id": log.user_id,
        "user_name": log.user.full_name if log.user else None,
        "action": log.action,
        "quantity_change": log.quantity_change,
        "previous_value": log.previous_value,
        "new_value": log.new_value,
        "note": log.note,
        "created_at": to_tr_time(log.created_at),
    } for log in logs]


@router.get("/export/excel")
def export_excel(
    category: Optional[str] = None,
    warehouse_id: Optional[int] = None,
    status: Optional[str] = None,
    low_stock: Optional[bool] = None,
    course_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_report_user),
):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    products = _report_products(db, current_user, category, warehouse_id, status, low_stock, course_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Ürün Raporu"

    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    critical_fill = PatternFill(start_color="FDE2E2", end_color="FDE2E2", fill_type="solid")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["No", "Ürün Adı", "Açıklama", "Kategori", "Dersler", "Depo", "Konum (Dolap / Raf / Bölme)",
               "Mevcut Stok", "Kritik Stok", "İdeal Stok", "Eksik (İdeale göre)", "Durum", "Son Güncelleme"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font, cell.border = header_fill, header_font, border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for n, p in enumerate(products, 1):
        stock = p.current_stock or 0
        values = [
            n, p.name, p.description or "", p.category,
            ", ".join(sorted(c.name for c in p.courses)),
            p.warehouse.name if p.warehouse else "", _location(p),
            stock, p.critical_stock, p.ideal_stock, max((p.ideal_stock or 0) - stock, 0), p.status,
            p.updated_at.replace(tzinfo=timezone.utc).astimezone(TZ_TR).strftime("%d.%m.%Y %H:%M") if p.updated_at else "",
        ]
        ws.append(values)
        is_critical = stock <= (p.critical_stock or 0)
        for cell in ws[ws.max_row]:
            cell.border = border
            cell.alignment = Alignment(vertical="center")
            if is_critical:
                cell.fill = critical_fill

    widths = [6, 38, 34, 24, 34, 24, 28, 12, 12, 12, 14, 22, 17]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"

    info = wb.create_sheet("Rapor Bilgisi")
    info.append([_title(db)])
    info.append(["Rapor tarihi", _tr_now().strftime("%d.%m.%Y %H:%M")])
    info.append(["Hazırlayan", current_user.full_name])
    info.append(["Ürün sayısı", len(products)])
    info.append(["Kritik stoktaki ürün", sum(1 for p in products if (p.current_stock or 0) <= (p.critical_stock or 0))])
    info.append(["Not", "Kırmızı satırlar kritik stok seviyesindeki ürünlerdir."])
    info["A1"].font = Font(bold=True, size=14)
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 50

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"deneyap_stok_raporu_{_tr_now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
_FONT_CANDIDATES = [
    (os.path.join(ASSETS_DIR, "fonts", "DejaVuSans.ttf"), os.path.join(ASSETS_DIR, "fonts", "DejaVuSans-Bold.ttf")),
    (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
]
_font_names = None


def _pdf_fonts():
    """Türkçe karakterleri (ş, ğ, ı, İ) destekleyen bir TTF yazı tipi kaydeder."""
    global _font_names
    if _font_names:
        return _font_names
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for regular, bold in _FONT_CANDIDATES:
        if os.path.exists(regular):
            pdfmetrics.registerFont(TTFont("TRFont", regular))
            pdfmetrics.registerFont(TTFont("TRFont-Bold", bold if os.path.exists(bold) else regular))
            from reportlab.lib.fonts import addMapping
            addMapping("TRFont", 0, 0, "TRFont")
            addMapping("TRFont", 1, 0, "TRFont-Bold")
            _font_names = ("TRFont", "TRFont-Bold")
            return _font_names
    _font_names = ("Helvetica", "Helvetica-Bold")  # son çare (Türkçe harfler bozulabilir)
    return _font_names


@router.get("/export/pdf")
def export_pdf(
    category: Optional[str] = None,
    warehouse_id: Optional[int] = None,
    status: Optional[str] = None,
    low_stock: Optional[bool] = None,
    course_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_report_user),
):
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    regular, bold = _pdf_fonts()
    products = _report_products(db, current_user, category, warehouse_id, status, low_stock, course_id)
    now = _tr_now()

    title_style = ParagraphStyle("t", fontName=bold, fontSize=16, leading=20, spaceAfter=4)
    normal = ParagraphStyle("n", fontName=regular, fontSize=9, leading=12)
    cell = ParagraphStyle("c", fontName=regular, fontSize=8, leading=10)
    head = ParagraphStyle("h", fontName=bold, fontSize=8.5, leading=10, textColor=colors.white)

    filters = []
    if warehouse_id:
        filters.append("Depo filtresi uygulandı")
    if category and category != "Tümü":
        filters.append(f"Kategori: {category}")
    if status:
        filters.append(f"Durum: {status}")
    if low_stock:
        filters.append("Yalnızca kritik stok")
    critical_count = sum(1 for p in products if (p.current_stock or 0) <= (p.critical_stock or 0))

    elements = [
        Paragraph(escape(_title(db)), title_style),
        Paragraph(escape(f"Rapor tarihi: {now.strftime('%d.%m.%Y %H:%M')}  •  Hazırlayan: {current_user.full_name}  •  "
                         f"Ürün sayısı: {len(products)}  •  Kritik stok: {critical_count}"), normal),
    ]
    if filters:
        elements.append(Paragraph(escape("Filtre: " + ", ".join(filters)), normal))
    elements.append(Spacer(1, 5 * mm))

    header = ["No", "Ürün Adı", "Kategori", "Depo / Konum", "Stok", "Kritik", "İdeal", "Durum"]
    data = [[Paragraph(h, head) for h in header]]
    critical_rows = []
    for i, p in enumerate(products, 1):
        loc = " / ".join(x for x in [p.warehouse.name if p.warehouse else "-", _location(p)] if x)
        data.append([
            Paragraph(str(i), cell), Paragraph(escape(p.name), cell), Paragraph(escape(p.category or ""), cell),
            Paragraph(escape(loc), cell), Paragraph(str(p.current_stock or 0), cell),
            Paragraph(str(p.critical_stock or 0), cell), Paragraph(str(p.ideal_stock or 0), cell),
            Paragraph(escape(p.status or ""), cell),
        ])
        if (p.current_stock or 0) <= (p.critical_stock or 0):
            critical_rows.append(i)

    if len(data) == 1:
        data.append([Paragraph("", cell), Paragraph("Bu filtreyle eşleşen ürün yok.", cell)] + [Paragraph("", cell)] * 6)

    table = Table(data, repeatRows=1, colWidths=[12 * mm, 77 * mm, 38 * mm, 58 * mm, 16 * mm, 16 * mm, 16 * mm, 34 * mm])
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A5F")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F4F8")]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9CA3AF")),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for r in critical_rows:
        style.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#FDE2E2")))
    table.setStyle(TableStyle(style))
    elements.append(table)

    def _footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(regular, 8)
        canvas.setFillColor(colors.HexColor("#6B7280"))
        canvas.drawString(15 * mm, 8 * mm, "Deneyap Depo ve Stok Yönetim Sistemi  •  Kırmızı satırlar kritik stoktadır")
        canvas.drawRightString(landscape(A4)[0] - 15 * mm, 8 * mm, f"Sayfa {doc.page}")
        canvas.restoreState()

    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), topMargin=14 * mm, bottomMargin=15 * mm,
                            leftMargin=15 * mm, rightMargin=15 * mm, title=_title(db), author=current_user.full_name)
    doc.build(elements, onFirstPage=_footer, onLaterPages=_footer)
    output.seek(0)
    filename = f"deneyap_stok_raporu_{now.strftime('%Y%m%d_%H%M')}.pdf"
    return StreamingResponse(output, media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename={filename}"})
